from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = "/Users/chltmddn5843/Downloads/skala/8주차/메디브릿지_MSA_프로젝트_기획안_보완본.docx"
NAVY = "17324D"
TEAL = "178C8C"
LIGHT = "EAF5F4"
PALE = "F3F6F8"
MID = "D6E1E8"
TEXT = "24323D"
MUTED = "657784"
WHITE = "FFFFFF"
RED = "B33A3A"
FONT = "Malgun Gothic"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.78)
sec.left_margin = sec.right_margin = Inches(0.82)
sec.header_distance, sec.footer_distance = Inches(0.35), Inches(0.35)

def font(run, size=10.5, bold=False, color=TEXT, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size, color, before, after in [
    ("Title", 28, NAVY, 0, 8),
    ("Subtitle", 13, MUTED, 0, 8),
    ("Heading 1", 17, NAVY, 14, 8),
    ("Heading 2", 13, TEAL, 11, 5),
    ("Heading 3", 11.5, NAVY, 8, 4),
]:
    s = styles[name]
    s.font.name = FONT
    s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    s.font.size = Pt(size)
    s.font.bold = name != "Subtitle"
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    if name == "Heading 1":
        s.paragraph_format.page_break_before = True

for style_name in ["List Bullet", "List Number"]:
    s = styles[style_name]
    s.font.name = FONT
    s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    s.font.size = Pt(10.5)
    s.paragraph_format.left_indent = Inches(0.28)
    s.paragraph_format.first_line_indent = Inches(-0.18)
    s.paragraph_format.space_after = Pt(4)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=90, start=120, bottom=90, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for tag, val in [("top",top),("start",start),("bottom",bottom),("end",end)]:
        el = tcMar.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            tcMar.append(el)
        el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa")

def set_cell_text(cell, text, bold=False, color=TEXT, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    font(p.add_run(str(text)), size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)

def table(headers, rows, widths=None, font_size=9.0):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(tbl_layout)
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=WHITE, size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(t.rows[0].cells[i], NAVY)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            if ri % 2 == 1: shade(cells[i], PALE)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
        tblPr = t._tbl.tblPr
        tblW = tblPr.find(qn("w:tblW"))
        tblW.set(qn("w:w"), str(sum(round(w*1440) for w in widths))); tblW.set(qn("w:type"), "dxa")
        grid = t._tbl.tblGrid
        for child in list(grid): grid.remove(child)
        for w in widths:
            gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(round(w*1440))); grid.append(gc)
        ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), "120"); ind.set(qn("w:type"), "dxa"); tblPr.append(ind)
        for row in t.rows:
            for i, w in enumerate(widths):
                tcW = row.cells[i]._tc.get_or_add_tcPr().find(qn("w:tcW"))
                tcW.set(qn("w:w"), str(round(w*1440))); tcW.set(qn("w:type"), "dxa")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def para(text="", bold=False, color=TEXT, size=10.5, align=None, after=6, italic=False):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    font(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p

def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(text), size=10.3)
    return p

def callout(title, body, fill=LIGHT):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    c = t.cell(0,0); c.width = Inches(6.5); shade(c, fill); margins(c, 150, 180, 150, 180)
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW")); tblW.set(qn("w:w"), "9360"); tblW.set(qn("w:type"), "dxa")
    ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), "180"); ind.set(qn("w:type"), "dxa"); tblPr.append(ind)
    grid = t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), "9360"); grid.append(gc)
    tcW = c._tc.get_or_add_tcPr().find(qn("w:tcW")); tcW.set(qn("w:w"), "9360"); tcW.set(qn("w:type"), "dxa")
    c.text = ""
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(4)
    font(p.add_run(title), 11, True, TEAL)
    p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.line_spacing = 1.12
    font(p2.add_run(body), 10.4, False, TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def page_break():
    pass

# Running header/footer
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(hp.add_run("MEDIBRIDGE  |  Agile & MSA Project"), 8.5, True, MUTED)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run("판교 4반 2조"), 8.5, False, MUTED)

# Cover
para("AGILE & MSA PROJECT", True, TEAL, 11, WD_ALIGN_PARAGRAPH.CENTER, 28)
para("메디브릿지", True, NAVY, 30, WD_ALIGN_PARAGRAPH.CENTER, 8)
para("B2B 의약품 공급·발주 관리 플랫폼", False, MUTED, 15, WD_ALIGN_PARAGRAPH.CENTER, 34)
callout("PROJECT PURPOSE", "CSO와 약국·병원을 연결해 비대면 의약품 탐색·추천·발주를 지원하고, 영업 인력의 지역적 한계를 보완해 더 많은 거래처에 제품을 제안한다.")
para("영업의 거리를 넘어, 의약품 공급과 발주를 연결하다.", True, NAVY, 14, WD_ALIGN_PARAGRAPH.CENTER, 48)
table(["역할", "담당"], [
    ("Product Owner", "김민주"), ("Scrum Master", "최승우"),
    ("Frontend", "임형준"), ("Backend", "권주현"), ("Service Planning", "전혜민")
], [2.0, 4.5], 9.5)
para("판교 4반 2조", False, MUTED, 10, WD_ALIGN_PARAGRAPH.CENTER, 0)

page_break()
doc.add_heading("1. 프로젝트 개요", level=1)
callout("한 줄 정의", "약국·병원 담당자가 의약품을 조회·발주하고, 결제 완료 후 발주를 확정하며 구매 이력에 따라 관련 품목을 추천받는 B2B 플랫폼")
table(["구분", "정의"], [
    ("Pain Point", "영업 인력과 지역에 의존하는 제품 제안 구조로 인해 거래처 확대와 제품 탐색에 한계가 있다."),
    ("서비스 기능", "거래처가 스스로 의약품을 탐색·추천받고 발주할 수 있는 비대면 업무 흐름을 제공한다."),
    ("서비스 가치", "영업의 거리를 넘어, 의약품 공급과 발주를 연결한다."),
], [1.35, 5.15], 9.3)
doc.add_heading("1.1 추진 배경", level=2)
para("CSO는 여러 제약사로부터 의약품 영업·판촉을 위탁받아 병원과 약국을 대상으로 제품 정보를 제공한다. 그러나 영업사원이 직접 방문할 수 있는 지역과 거래처 수에는 한계가 있고, 거래처를 늘리려면 영업 인력과 비용도 함께 증가한다.")
para("또한 유사한 성분과 효능의 제품이 여러 제약사에서 판매되기 때문에 구매 담당자가 자신의 구매 특성과 수요에 맞는 제품을 직접 비교하기 어렵다. 거래처별 주문 이력과 제품별 수요 데이터 역시 제품 제안 과정에 충분히 활용되지 못하고 있다.")
doc.add_heading("1.2 서비스 목표", level=2)
for x in [
    "지역과 관계없이 의약품을 탐색하고 발주할 수 있는 비대면 채널 제공",
    "거래처별 구매 이력을 활용한 설명 가능한 제품 추천",
    "결제와 발주 확정을 서비스 간 통신으로 자동 처리",
    "CSO 운영사가 적은 영업 인력으로 더 많은 거래처를 관리할 수 있도록 지원",
]: bullet(x)
doc.add_heading("1.3 기대 효과", level=2)
callout("핵심 효과", "비대면 의약품 탐색·추천·발주를 통해 지역 간 영업 격차를 줄이고, 제한된 영업 인력으로 더 많은 거래처와 제약회사를 연결한다.", PALE)

page_break()
doc.add_heading("2. 이해관계자와 Pain Point", level=1)
table(["이해관계자", "현재 불편", "제공 가치"], [
    ("약국·병원", "지역과 영업 접점에 따라 제품 정보가 제한되고, 유사 제품 비교가 어렵다.", "지역과 관계없이 다양한 의약품을 탐색·추천받고 편리하게 발주"),
    ("CSO 기업", "거래처 확대가 영업 인력과 비용 증가로 이어지고, 제안이 개인 경험과 기존 관계에 의존한다.", "적은 영업 인력으로 더 많은 지역과 거래처에 데이터 기반 제품 제안"),
    ("제약회사", "기존 영업망이 닿지 않는 거래처에는 제품을 노출하기 어렵다.", "제품 노출 범위와 정상적인 판매 기회 확대"),
], [1.2, 2.7, 2.6], 9.1)
doc.add_heading("2.1 공통 핵심 문제", level=2)
callout("PAIN POINT", "거래처별 구매 이력과 제품별 수요 데이터가 영업·구매 의사결정에 충분히 활용되지 못하고 있다.")
doc.add_heading("2.2 도입 의사결정 관점", level=2)
para("플랫폼 도입의 핵심 의사결정자는 개별 영업사원이 아니라 CSO 운영사다. 메디브릿지는 영업사원을 전면 대체하기보다 반복적인 제품 안내와 발주 업무를 자동화하여 영업 커버리지를 확대한다.")
para("특정 제품을 임의로 우대하는 방식보다 실제 구매 데이터를 기반으로 적합한 제품을 추천함으로써 구매전환율·재구매율·고객 유지율을 높이고 컴플라이언스 리스크를 줄이는 것을 지향한다.")

page_break()
doc.add_heading("3. AI 솔루션", level=1)
doc.add_heading("3.1 핵심 사용자 흐름", level=2)
table(["단계", "사용자 행동", "시스템 처리"], [
    ("1", "계정 등록 및 로그인", "역할에 맞는 접근 권한 확인"),
    ("2", "의약품 탐색", "약효군별 목록·상세 정보 제공"),
    ("3", "발주 신청", "제품 존재와 중복 신청 검증 후 PENDING 생성"),
    ("4", "결제 처리", "거래 UUID 생성 및 payment.completed 발행"),
    ("5", "발주 확정", "이벤트 수신 후 주문 상태를 ACTIVE로 변경"),
    ("6", "추천 확인", "구매 이력 기반 미구매 연관 품목 추천"),
], [0.55, 2.1, 3.85], 9.2)
doc.add_heading("3.2 AI의 역할", level=2)
callout("MVP 추천 규칙", "구매 이력이 있으면 가장 많이 구매한 약효군을 찾고, 해당 약효군에서 아직 구매하지 않은 제품을 추천한다. 구매 이력이 없는 신규 거래처에는 누적 발주 건수가 높은 인기 제품을 추천한다.")
doc.add_heading("3.3 발전 방향", level=2)
for x in [
    "거래처 규모·진료과가 유사한 기관의 구매 패턴을 활용한 추천",
    "주문 증가율·재주문율·동일 계열 등 추천 이유 제공",
    "구매 주기를 활용한 예상 재주문 시점 알림",
    "성분·함량·제형·허가사항을 고려한 대체 제품 추천",
]: bullet(x)
para("※ 확장 기능은 추가 데이터가 필요한 후속 백로그이며, 이번 실습의 구현 범위에는 포함하지 않는다.", False, MUTED, 9.3, after=0, italic=True)

page_break()
doc.add_heading("4. Agile 실행 계획", level=1)
doc.add_heading("4.1 백로그 우선순위 기준", level=2)
table(["기준", "판단 질문"], [
    ("사용자 가치", "의약품 탐색과 발주라는 핵심 가치를 제공하는 데 반드시 필요한가?"),
    ("기능 의존성", "다음 기능이 실행되기 전에 먼저 생성되어야 하는 데이터인가?"),
    ("기술 위험", "REST·Kafka·추천 연동을 단계적으로 검증할 수 있는가?"),
    ("실현 가능성", "제공된 템플릿과 실습 시간 안에 동작 결과를 보여줄 수 있는가?"),
], [1.35, 5.15], 9.3)
doc.add_heading("4.2 Sprint 구분", level=2)
table(["구분", "목표", "포함 기능", "분리 이유"], [
    ("Sprint 1", "발주 가능한 최소 서비스", "회원가입·로그인, 의약품 등록·조회, 발주 신청·조회", "계정 → 의약품 → 발주의 완결된 Walking Skeleton을 먼저 검증"),
    ("Sprint 2", "거래 자동화와 추천 확장", "결제, Kafka 상태 변경, 구매 이력 기반 추천", "Sprint 1에서 생성된 데이터와 피드백을 바탕으로 연동 위험이 큰 기능을 확장"),
], [0.9, 1.45, 2.0, 2.15], 8.7)
callout("상태 정의", "PENDING은 결제 대기, ACTIVE는 결제 완료에 따른 발주 확정, CANCELLED는 발주 취소를 의미한다. 실제 입고 완료는 배송·입고 이벤트가 없으므로 이번 범위에서 사용하지 않는다.", PALE)
doc.add_heading("4.3 이해관계자 피드백 반영", level=2)
table(["구분", "내용"], [
    ("초기 가설", "추천 플랫폼의 핵심 사용자를 CSO 영업사원으로 설정"),
    ("피드백", "영업사원이 추천 결과를 수용할 유인이 충분한지 검토 필요"),
    ("변경", "도입 의사결정자를 CSO 운영사로 재정의하고, 영업 대체보다 커버리지 확대를 핵심 가치로 설정"),
    ("백로그 영향", "셀프 탐색·발주를 Sprint 1의 핵심 가치로, 자동 결제·추천을 Sprint 2 확장 기능으로 배치"),
], [1.35, 5.15], 9.1)

page_break()
doc.add_heading("5. MSA 설계", level=1)
doc.add_heading("5.1 서비스 경계", level=2)
table(["현재 코드명", "업무상 명칭", "책임"], [
    ("user-service", "partner-service", "CSO 및 약국·병원 사용자·거래처 관리"),
    ("course-service", "product-service", "의약품 등록·분류·제품 정보 관리"),
    ("enrollment-service", "order-service", "의약품 발주 신청과 상태 관리"),
    ("payment-service", "payment-service", "구매 승인, 결제와 거래 기록 관리"),
    ("recommend-service", "recommend-service", "거래처 구매 이력 기반 의약품 추천"),
], [1.45, 1.5, 3.55], 9.2)
doc.add_heading("5.2 서비스 간 통신", level=2)
callout("통신 흐름", "Client → API Gateway → Partner / Product / Order / Recommend\nOrder → Product (제품 확인, REST)\nOrder → Payment (결제 요청, REST)\nPayment → Kafka: payment.completed → Order 상태 ACTIVE\nOrder → Kafka: enrollment.completed → Recommend 갱신", LIGHT)
doc.add_heading("5.3 분할 원칙과 인프라", level=2)
para("서비스는 업무 책임과 데이터의 생성·변경 주체를 기준으로 분리한다. 각 서비스는 자신이 소유한 데이터를 직접 변경하고, 다른 서비스의 데이터는 REST API 또는 Kafka 이벤트를 통해 사용한다.")
para("실습에서는 API 계약과 인프라 설정 변경을 최소화하기 위해 실제 디렉터리·경로는 기존 코드명을 유지하고, 화면과 보고서에서 업무상 명칭으로 치환한다.", False, MUTED, 9.3, after=6, italic=True)
table(["인프라", "역할"], [
    ("API Gateway", "클라이언트 요청의 단일 진입점과 서비스별 라우팅"),
    ("Auth Server", "OAuth2 로그인과 토큰 발급"),
    ("Eureka", "서비스 등록과 탐색"),
    ("Kafka", "결제 완료 및 발주 확정 이벤트 전달"),
    ("MariaDB", "서비스별 업무 데이터 저장"),
], [1.5, 5.0], 9.2)

page_break()
doc.add_heading("6. 데이터 설계", level=1)
doc.add_heading("6.1 도메인 매핑", level=2)
table(["기존 데이터", "업무 의미", "핵심 필드"], [
    ("users", "거래처·담당자", "id, email, password, name, role"),
    ("courses", "의약품", "id, title, description, category, price, instructor_id, enrollment_count, status"),
    ("enrollments", "발주", "id, user_id, course_id, status, created_at, updated_at"),
    ("payments", "결제·정산", "id, user_id, course_id, amount, status, transaction_id"),
], [1.25, 1.55, 3.7], 8.8)
doc.add_heading("6.2 내부 코드값의 업무 의미", level=2)
table(["구분", "코드값", "업무 의미"], [
    ("역할", "INSTRUCTOR", "제약사·CSO 관리자"),
    ("역할", "STUDENT", "약국·병원 구매 담당자"),
    ("발주", "PENDING", "결제 대기"),
    ("발주", "ACTIVE", "발주 확정"),
    ("발주", "CANCELLED", "발주 취소"),
    ("결제", "COMPLETED", "결제 완료"),
    ("결제", "FAILED", "결제 실패"),
], [1.0, 1.5, 4.0], 9.2)
doc.add_heading("6.3 주요 관계", level=2)
for x in [
    "CSO 관리자는 여러 의약품을 등록한다.",
    "약국·병원 구매 담당자는 여러 의약품을 발주한다.",
    "제품별 발주 이력과 거래처별 결제 이력을 관리한다.",
    "추천 서비스는 발주·제품 데이터를 읽어 추천 결과를 계산한다.",
]: bullet(x)

page_break()
doc.add_heading("7. Product Backlog", level=1)
stories = [
    ("US-01", "계정", "CSO·구매 담당자가 이메일·비밀번호·역할로 계정을 등록한다.", "이메일 중복 검증, BCrypt 암호화, 역할 분리", "S1"),
    ("US-02", "계정", "거래처 담당자가 토큰 기반으로 내 정보를 조회한다.", "X-User-Id 기반 사용자 조회", "S1"),
    ("US-03", "제품", "CSO 관리자가 신규 의약품을 등록한다.", "권한 확인, ACTIVE 등록, 누적 발주 0", "S1"),
    ("US-04", "제품", "구매 담당자가 약효군별 의약품을 조회한다.", "분류 필터, 상세 정보, 누적 발주 노출", "S1"),
    ("US-06", "발주", "구매 담당자가 의약품 발주를 신청한다.", "제품 확인, 중복 방지, PENDING 생성", "S1"),
    ("US-07", "발주", "구매 담당자가 내 발주 내역과 상태를 조회한다.", "본인 발주, 제품 정보, 상태 표시", "S1"),
    ("US-05", "제품", "구매 확정 시 누적 발주 건수를 증가시킨다.", "내부 REST, 원자적 카운트 증가", "S2"),
    ("US-08", "발주", "결제 완료 이벤트로 발주를 확정한다.", "Kafka 수신, ACTIVE 변경, 후속 이벤트", "S2"),
    ("US-09", "결제", "내부 결제를 처리하고 거래 UUID를 발급한다.", "COMPLETED, UUID, 이벤트 발행", "S2"),
    ("US-10", "결제", "거래처별 결제 내역과 거래번호를 조회한다.", "결제 목록과 증빙 정보 제공", "S2"),
    ("US-11", "추천", "구매 이력 기반 미구매 연관 제품을 추천한다.", "최빈 약효군, 미구매 제품, 상위 5개", "S2"),
    ("US-12", "추천", "신규 거래처에 인기 의약품을 추천한다.", "이력 없음 식별, 누적 발주순 상위 5개", "S2"),
]
table(["ID", "Epic", "사용자 스토리", "수용 기준 핵심", "Sprint"], stories, [0.62, 0.72, 2.65, 2.05, 0.46], 7.9)

page_break()
doc.add_heading("8. 검증 포인트와 후속 과제", level=1)
doc.add_heading("8.1 이번 실습의 완료 기준", level=2)
for x in [
    "사용자가 로그인한 뒤 의약품을 조회하고 발주할 수 있다.",
    "발주 생성 시 PENDING 상태가 확인된다.",
    "결제 완료 이벤트 수신 후 발주 상태가 ACTIVE로 변경된다.",
    "구매 이력 유무에 따라 다른 추천 결과가 반환된다.",
    "Swagger UI와 화면 캡처로 요청 전·후 상태 변화를 증빙한다.",
]: bullet(x)
doc.add_heading("8.2 알려진 제약", level=2)
table(["제약", "현재 처리", "후속 개선"], [
    ("결제 금액", "템플릿의 고정 금액 사용 가능", "제품 단가와 수량을 결제 요청에 반영"),
    ("재발주", "동일 사용자·제품 중복 제한", "PENDING만 중복 차단하고 완료 후 재발주 허용"),
    ("입고 상태", "결제 완료까지만 확인", "배송·입고 이벤트와 별도 상태 추가"),
    ("고급 추천", "약효군·누적 발주 기반 규칙 추천", "거래처 특성·재고·주기·대체 성분 데이터 추가"),
], [1.25, 2.35, 2.9], 9.0)
doc.add_heading("8.3 결론", level=2)
callout("PROJECT OUTCOME", "메디브릿지는 이해관계자의 요구를 핵심 발주 흐름과 확장 기능으로 나누어 두 번의 Sprint로 검증한다. 업무와 데이터 책임에 따라 서비스를 분리하고, REST와 Kafka를 통해 독립된 서비스가 하나의 거래 흐름을 완성하도록 설계했다.")

# Keep tables from splitting rows; set metadata.
for t in doc.tables:
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit"); trPr.append(cant)

doc.core_properties.title = "메디브릿지 MSA 프로젝트 기획안 보완본"
doc.core_properties.subject = "Agile & MSA 조별 실습 기획안"
doc.core_properties.author = "판교 4반 2조"
doc.save(OUT)
print(OUT)
