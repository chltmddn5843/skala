from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_ORIENT

ROOT = "/Users/chltmddn5843/Downloads/skala/8주차"
OUT = f"{ROOT}/Agile+MSA_과제_4반_P131_최승우.docx"
ARCH = f"{ROOT}/학습자료/Service Container Flow-2026-09-02-050713.png"
NOTE = f"{ROOT}/Day1/image/서브노트/1788325666636.png"

NAVY, TEAL, LIGHT, PALE = "17324D", "178C8C", "EAF5F4", "F3F6F8"
TEXT, MUTED, WHITE, RED = "24323D", "657784", "FFFFFF", "A13A3A"
FONT = "Malgun Gothic"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.78)
sec.left_margin = sec.right_margin = Inches(0.82)
sec.header_distance = sec.footer_distance = Inches(0.35)

def set_font(run, size=10.5, bold=False, color=TEXT, italic=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), FONT)
    rpr.rFonts.set(qn("w:hAnsi"), FONT)
    rpr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size, run.bold, run.italic = Pt(size), bold, italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run

normal = doc.styles["Normal"]
normal.font.name, normal.font.size = FONT, Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size, color, before, after in [
    ("Title", 28, NAVY, 0, 8), ("Subtitle", 13, MUTED, 0, 8),
    ("Heading 1", 17, NAVY, 14, 8), ("Heading 2", 13, TEAL, 11, 5),
    ("Heading 3", 11.5, NAVY, 8, 4),
]:
    s = doc.styles[name]
    s.font.name, s.font.size = FONT, Pt(size)
    s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    s.font.bold = name != "Subtitle"
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
    s.paragraph_format.keep_with_next = True
for style_name in ["List Bullet", "List Number"]:
    s = doc.styles[style_name]
    s.font.name, s.font.size = FONT, Pt(10.3)
    s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    s.paragraph_format.left_indent, s.paragraph_format.first_line_indent = Inches(.32), Inches(-.18)
    s.paragraph_format.space_after, s.paragraph_format.line_spacing = Pt(4), 1.12

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd")) or OxmlElement("w:shd")
    if shd.getparent() is None: tcpr.append(shd)
    shd.set(qn("w:fill"), fill)

def cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar"); tcpr.append(mar)
    for tag, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        el = mar.find(qn(f"w:{tag}"))
        if el is None: el = OxmlElement(f"w:{tag}"); mar.append(el)
        el.set(qn("w:w"), str(value)); el.set(qn("w:type"), "dxa")

def set_cell(cell, value, bold=False, color=TEXT, size=9.0, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]; p.alignment = align
    p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(0), 1.06
    set_font(p.add_run(str(value)), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)

def table(headers, rows, widths, size=8.9):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(layout)
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, WHITE, 9.0, WD_ALIGN_PARAGRAPH.CENTER)
        shade(t.rows[0].cells[i], NAVY)
    hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); t.rows[0]._tr.get_or_add_trPr().append(hdr)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value, size=size, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            if ri % 2: shade(cells[i], PALE)
    dxa = [round(w * 1440) for w in widths]
    tblpr = t._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW")); tblw.set(qn("w:w"), str(sum(dxa))); tblw.set(qn("w:type"), "dxa")
    ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), "120"); ind.set(qn("w:type"), "dxa"); tblpr.append(ind)
    grid = t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in dxa:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(width)); grid.append(gc)
    for row in t.rows:
        cant = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(cant)
        for i, width in enumerate(dxa):
            tcw = row.cells[i]._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tcw.set(qn("w:w"), str(width)); tcw.set(qn("w:type"), "dxa")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def para(text="", bold=False, color=TEXT, size=10.5, align=None, after=6, italic=False):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text), size, bold, color, italic)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    set_font(p.add_run(text), 10.3)
    return p

def callout(title, body, fill=LIGHT):
    t = doc.add_table(rows=1, cols=1); t.autofit = False
    c = t.cell(0, 0); shade(c, fill); cell_margins(c, 150, 180, 150, 180)
    tblpr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblpr.append(layout)
    tblw = tblpr.find(qn("w:tblW")); tblw.set(qn("w:w"), "9360"); tblw.set(qn("w:type"), "dxa")
    ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), "180"); ind.set(qn("w:type"), "dxa"); tblpr.append(ind)
    grid = t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), "9360"); grid.append(gc)
    tcw = c._tc.get_or_add_tcPr().find(qn("w:tcW")); tcw.set(qn("w:w"), "9360"); tcw.set(qn("w:type"), "dxa")
    c.text = ""; p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(title), 11, True, TEAL)
    p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.line_spacing = 1.12
    set_font(p2.add_run(body), 10.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def page_break(): doc.add_page_break()

def caption(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(8)
    set_font(p.add_run(text), 8.8, False, MUTED, True)

hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(hp.add_run("AGILE & MSA  |  개인과제"), 8.5, True, MUTED)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(fp.add_run("판교 4반 P131 최승우"), 8.5, False, MUTED)

# Cover
para("AGILE & MSA REPORT", True, TEAL, 11, WD_ALIGN_PARAGRAPH.CENTER, 28)
para("메디브릿지", True, NAVY, 30, WD_ALIGN_PARAGRAPH.CENTER, 8)
para("B2B 의약품 공급·발주 관리 AI 서비스", False, MUTED, 15, WD_ALIGN_PARAGRAPH.CENTER, 34)
callout("REPORT PURPOSE", "이해관계자의 가치를 유지하면서 AI 서비스를 빠르게 검증하기 위해 MVP, Agile, MSA를 어떻게 결합할지 설명하고, 실습 아키텍처와 시행착오를 근거로 제안한다.")
para("업무의 신뢰를 지키면서, 변화는 더 빠르고 작게.", True, NAVY, 14, WD_ALIGN_PARAGRAPH.CENTER, 46)
table(["구분", "내용"], [
    ("과정", "Cloud Agile 방법론 및 MSA 개발"),
    ("소속", "판교 4반"), ("학번", "P131"), ("작성자", "최승우"),
    ("대상 서비스", "메디브릿지 - CSO·약국·병원 B2B 플랫폼"),
], [1.5, 5.0], 9.5)

page_break()
doc.add_heading("보고서 요약", level=1)
callout("핵심 결론", "메디브릿지의 MVP는 '의약품 탐색 → 발주 → 결제 완료 → 발주 확정 → 구매 이력 기반 추천'의 최소 흐름이다. Agile은 현장 피드백으로 우선순위를 조정하고, MSA는 업무·데이터 책임을 분리해 추천이나 결제 기능을 바꿔도 핵심 발주 흐름의 영향을 줄인다.")
doc.add_heading("문제-해결-검증의 연결", level=2)
table(["관점", "핵심 질문", "본 보고서의 답"], [
    ("가치", "누가 어떤 불편을 겪는가?", "약국·병원의 정보 격차, CSO의 영업 확장 비용, 제약사의 제한된 노출"),
    ("솔루션", "무엇을 먼저 만들 것인가?", "설명 가능한 규칙 기반 추천을 포함한 발주 가능한 MVP"),
    ("구조", "어떻게 자주 바꿀 것인가?", "업무 단위 Microservice, API Gateway, Eureka, Kafka, 서비스별 데이터 책임"),
    ("검증", "무엇으로 성공을 판단할 것인가?", "발주 완료율, 추천 클릭·전환율, 처리 시간, 변경 리드타임"),
], [1.0, 2.1, 3.4], 9.1)
doc.add_heading("적용 범위", level=2)
for x in [
    "이번 MVP는 의약품 정보 제공과 발주 지원을 목적으로 하며, 의료 진단·처방 판단을 대신하지 않는다.",
    "고급 개인화 모델은 데이터와 검증 근거가 쌓인 뒤 후속 Sprint에서 도입한다.",
    "특정 제품을 임의 우대하지 않고 추천 이유를 함께 제시해 신뢰와 컴플라이언스를 보호한다.",
]: bullet(x)

page_break()
doc.add_heading("1. 이해관계자 가치 (Pain Point)", level=1)
doc.add_heading("1.1 MVP가 필요한 경우", level=2)
para("현업의 AI 도입은 완전히 새로운 서비스를 만드는 일보다 기존 업무 흐름에 AI 기능을 붙이는 경우가 많다. 이때 문제·데이터·사용자 반응이 충분히 검증되지 않은 상태에서 전체 기능을 한 번에 개발하면 비용은 커지고 잘못된 가정을 수정하기 어렵다. 따라서 핵심 가설을 가장 작은 업무 흐름으로 구현해 실제 사용자의 반응을 확인할 때 MVP가 적합하다.")
callout("메디브릿지의 MVP 가설", "구매 이력에 근거한 설명 가능한 추천과 비대면 발주 채널을 제공하면, CSO는 영업 인력을 같은 비율로 늘리지 않고도 거래처 접점을 확대할 수 있고 약국·병원은 지역과 담당자에 덜 의존해 제품을 탐색할 수 있다.")
doc.add_heading("1.2 이해관계자별 Pain Point와 가치", level=2)
table(["이해관계자", "Agile·MSA가 없을 때의 불편", "제공 가치"], [
    ("약국·병원", "지역·영업 접점에 따라 정보가 달라지고, 유사 제품 비교와 발주 확인이 번거롭다.", "언제 어디서나 탐색·발주하고 구매 이력에 맞는 제품과 추천 이유를 확인"),
    ("CSO 운영사", "거래처 확대가 인력·방문 비용 증가로 직결되고, 변경 시 전체 시스템 영향 때문에 개선이 늦다.", "반복 안내·발주 자동화, 데이터 기반 제안, 기능별 독립 변경으로 운영 효율 향상"),
    ("제약회사", "기존 영업망 밖 거래처에 제품을 알리기 어렵고 수요 데이터를 활용하기 어렵다.", "제품 노출 범위 확대와 실제 수요 신호 확보"),
    ("개발·운영팀", "기능이 얽히면 작은 수정도 전체 배포·장애 위험으로 이어진다.", "서비스별 책임·배포 범위 명확화, 장애 격리, 진행 상황 가시화"),
], [1.05, 2.85, 2.60], 8.7)
doc.add_heading("1.3 왜 Agile과 MSA가 함께 필요한가", level=2)
para("Agile은 짧은 Sprint마다 '지금 가장 가치 있는 변화'를 선택하고 피드백을 다음 백로그에 반영하는 운영 방식이다. MSA는 그 변화를 실제로 작게 배포할 수 있도록 업무와 데이터 책임을 분리하는 기술 구조다. 방법론만 있고 구조가 결합돼 있으면 변경 비용이 줄지 않고, 구조만 나눠 놓고 사용자 피드백이 없으면 서비스만 많아진다.")
callout("강의에서 얻은 기준", "이해관계자의 가치를 유지한 채 표면적인 서비스를 변화시킨다. 기능을 고칠 때 다른 서비스에 영향을 주지 않도록 독립성을 확보하되, 이해하지 못한 업무를 무리하게 나누지 않고 데이터의 생성·처리 책임을 기준으로 경계를 정한다.", PALE)

page_break()
doc.add_heading("2. 이를 해결하기 위한 AI 솔루션", level=1)
doc.add_heading("2.1 핵심 업무 흐름", level=2)
table(["단계", "사용자 행동", "서비스 처리", "검증 포인트"], [
    ("1", "계정 등록·로그인", "권한 확인 및 토큰 발급", "역할별 접근 분리"),
    ("2", "의약품 탐색", "약효군별 목록·상세 제공", "검색·조회 성공"),
    ("3", "발주 신청", "제품·중복 검증 후 PENDING 생성", "대기 상태 확인"),
    ("4", "결제 처리", "거래 UUID 생성, payment.completed 발행", "결제 기록 생성"),
    ("5", "발주 확정", "이벤트 수신 후 ACTIVE 변경", "비동기 상태 전환"),
    ("6", "추천 확인", "구매 이력 기반 미구매 품목 반환", "추천 이유·결과 확인"),
], [.45, 1.45, 2.75, 1.85], 8.3)
doc.add_heading("2.2 규칙 기반 AI MVP", level=2)
callout("추천 규칙", "구매 이력이 있으면 가장 많이 구매한 약효군을 찾고 해당 약효군에서 아직 구매하지 않은 제품을 최대 5개 추천한다. 이력이 없는 신규 거래처에는 누적 발주 건수가 높은 인기 제품을 추천한다.")
para("규칙 기반 모델은 결과를 설명하기 쉽고 초기 데이터가 적어도 동작하므로 MVP에 적합하다. 추천 문구에는 '자주 구매한 약효군의 미구매 제품' 또는 '신규 거래처 인기 제품'처럼 근거를 노출한다. 이는 10년간 쌓은 업무 신뢰가 부정확한 AI 결과로 훼손되는 위험을 줄인다.")
doc.add_heading("2.3 Agile과 MSA의 역할", level=2)
table(["구분", "담당 역할", "메디브릿지 적용"], [
    ("Agile", "가설·우선순위·피드백 관리", "Walking Skeleton을 먼저 검증하고 추천·결제 연동을 다음 Sprint로 확장"),
    ("MSA", "변경·배포·장애 영향 범위 축소", "Partner, Product, Order, Payment, Recommend를 업무 책임별 분리"),
    ("REST", "즉시 응답이 필요한 동기 조회", "제품 존재 확인, 결제 요청, 추천용 이력·제품 조회"),
    ("Kafka", "서비스 간 비동기 상태 전달", "payment.completed 및 enrollment.completed 이벤트 처리"),
], [1.0, 2.15, 3.35], 9.0)
doc.add_heading("2.4 두 번의 Sprint", level=2)
table(["Sprint", "목표", "기능", "피드백·판단"], [
    ("1", "발주 가능한 최소 서비스", "계정, 제품 조회, 발주 신청·조회", "사용자가 제품을 찾고 PENDING 발주를 만들 수 있는가?"),
    ("2", "거래 자동화와 추천", "결제, Kafka 상태 변경, 규칙 추천", "결제 후 ACTIVE 전환과 추천 근거가 이해되는가?"),
], [.7, 1.5, 2.0, 2.3], 8.8)
para("후속 백로그: 유사 기관 구매 패턴, 재주문 시점, 성분·함량·제형을 고려한 대체 제품 추천. 추가 데이터와 성능 기준이 확보될 때만 도입한다.", False, MUTED, 9.3, italic=True)

page_break()
doc.add_heading("3. 아키텍처 구성도", level=1)
doc.add_heading("3.1 전체 파이프라인", level=2)
para("사용자 요청은 Vue Frontend에서 API Gateway로 들어온다. Gateway는 Eureka에서 서비스 위치를 찾아 요청을 라우팅하고, Auth Server가 발급한 토큰으로 접근 권한을 확인한다. 업무 서비스는 REST로 즉시 필요한 정보를 조회하고 Kafka로 결제·발주 완료 이벤트를 전달한다. Recommend Service는 발주·제품 데이터를 읽어 규칙 기반 추천을 계산한다.")
doc.add_picture(ARCH, width=Inches(6.45))
doc.inline_shapes[-1]._inline.docPr.set("descr", "API Gateway, 마이크로서비스, Eureka, Kafka, MariaDB, Docker Compose 간 연결을 나타낸 실습 아키텍처 구성도")
caption("그림 1. 실습 코드의 Service Container Flow (강의 실습 자료)")

page_break()
doc.add_heading("3.2 구성요소와 책임", level=2)
table(["구성요소", "실습 서비스", "메디브릿지 업무 책임"], [
    ("Frontend", "vue-frontend", "약국·병원 및 CSO의 탐색·발주·관리 화면"),
    ("Gateway", "api-gateway :8080", "단일 진입점, 인증 연계, 경로 라우팅"),
    ("Partner", "user-service :8081", "거래처·담당자·역할 관리"),
    ("Product", "course-service :8082", "의약품·약효군·가격·노출 상태 관리"),
    ("Order", "enrollment-service :8083", "발주 생성과 PENDING/ACTIVE/CANCELLED 상태 관리"),
    ("Payment", "payment-service :8084", "결제·거래 UUID·완료 이벤트"),
    ("Recommend", "recommend-service :8085", "구매 이력 기반 규칙 추천"),
    ("Infra", "Eureka/Kafka/MariaDB/Auth", "서비스 탐색, 이벤트, 저장, OAuth2/JWT 인증"),
], [1.05, 2.0, 3.45], 8.7)
doc.add_heading("3.3 데이터와 통신 경계", level=2)
for x in [
    "서비스는 업무 결과로 생성되는 데이터의 소유권을 기준으로 나눈다. 예: Product가 제품을, Order가 발주 상태를 변경한다.",
    "다른 서비스의 DB를 직접 수정하지 않고 REST API나 Kafka 이벤트로 협업한다.",
    "Eureka는 서비스의 IP·포트를 하드코딩하지 않도록 이름 기반 탐색을 제공한다.",
    "Docker Compose는 컨테이너와 네트워크를 함께 실행하며 healthcheck·depends_on으로 기동 순서를 관리한다.",
]: bullet(x)
callout("User Story에서 배포까지", "Pain Point → Product Backlog → Sprint 선택 → 서비스별 API 계약 → 구현·컨테이너 실행 → REST/Kafka 통합 검증 → 사용자 피드백 → 다음 Sprint 백로그 갱신", PALE)

page_break()
doc.add_heading("4. 트러블 슈팅", level=1)
doc.add_heading("4.1 실습에서 확인한 문제와 해결", level=2)
table(["증상", "원인·판단", "해결 및 확인"], [
    ("서비스가 한꺼번에 정상 기동하지 않음", "DB·Eureka·Auth·Kafka가 준비되기 전에 업무 서비스가 시작될 수 있음", "healthcheck와 depends_on을 확인하고 MariaDB/Kafka → Eureka → Auth → Gateway·업무 서비스 순으로 로그 점검"),
    ("서비스 주소 변경 시 호출 실패 우려", "IP·포트를 직접 연결하면 컨테이너 재시작 시 주소가 달라질 수 있음", "각 서비스를 Eureka에 등록하고 Gateway와 서비스가 이름으로 탐색하도록 구성"),
    ("결제 후 발주 상태가 즉시 바뀌지 않음", "결제 완료는 Kafka 비동기 이벤트이므로 발행·소비 어느 한쪽이 실패할 수 있음", "payment.completed 발행 로그, broker 상태, Order consumer 로그, ACTIVE 변경을 순서대로 확인"),
    ("추천 서비스가 등록 또는 갱신되지 않음", "FastAPI 시작 시 Eureka 초기화나 Kafka Consumer 시작이 실패할 수 있음", "lifespan 로그에서 등록·Consumer 시작을 확인하고 환경변수 URL·host·port를 점검"),
    ("localhost 접속이 예상과 다름", "브라우저 HTTPS 자동 전환 또는 host 설정 영향", "http://localhost:포트를 명시하고 127.0.0.1 및 서비스 health endpoint와 비교"),
    ("Docker 이미지·실행 환경 불일치", "배포 이미지의 CPU 아키텍처나 태그가 로컬과 다를 수 있음", "docker images로 태그를 확인하고 제공된 ARM64 이미지·compose 파일의 platform 값을 사용"),
], [1.55, 2.35, 2.60], 8.15)
doc.add_heading("4.2 새로 이해한 용어", level=2)
table(["용어", "이해한 의미"], [
    ("Walking Skeleton", "사용자에게 보이는 화면부터 핵심 백엔드·DB까지 가장 얇게 연결한 실행 가능한 전체 흐름"),
    ("Service Discovery", "서비스가 자신의 위치를 등록하고 다른 서비스가 이름으로 찾아가는 방식. 실습에서는 Eureka 사용"),
    ("API Gateway", "클라이언트 요청을 한 곳에서 받고 인증·라우팅을 처리하는 단일 진입점"),
    ("Event-driven", "한 서비스의 업무 완료 사실을 이벤트로 발행하고 다른 서비스가 비동기로 반응하는 구조"),
    ("XaaS/PaaS", "직접 OS·인프라를 구성하는 부담을 줄이고 플랫폼 기능을 서비스로 제공받는 방식"),
], [1.45, 5.05], 9.0)
doc.add_heading("4.3 재현 가능한 점검 순서", level=2)
for x in [
    "docker images에서 필요한 이미지와 태그를 확인한다.",
    "docker compose up -d 후 docker compose ps로 health 상태를 확인한다.",
    "Eureka(:8761)에서 서비스 등록 여부를 확인한다.",
    "Gateway(:8080)를 통해 REST API를 호출하고 요청·응답 계약을 확인한다.",
    "결제 전후 발주 상태와 Kafka producer/consumer 로그를 함께 비교한다.",
    "구매 이력 유무에 따라 추천 결과와 추천 이유가 달라지는지 캡처한다.",
]: bullet(x)

page_break()
doc.add_heading("5. 레퍼런스 (필요한 스냅샷)", level=1)
doc.add_heading("5.1 강의 메모 핵심", level=2)
table(["키워드", "정리"], [
    ("이해관계자 가치", "AI 기능보다 먼저 누가 어떤 불편을 겪는지 정의하며, 조직의 가치와 신뢰를 유지한다."),
    ("Agile", "짧은 Sprint, 진행 상황 가시화, 현장 피드백 반영으로 점진적으로 발전시킨다."),
    ("MSA", "빠르고 잦은 변경을 가능하게 하며 업무·데이터 생성과 처리 책임을 기준으로 분할한다."),
    ("B2B AI", "기존 업무에 AI를 붙여 반복 업무와 의사결정을 지원하는 문제를 우선한다."),
    ("검증", "백엔드·프론트 동작, API 계약, 상태 변화, 아키텍처를 스냅샷으로 증명한다."),
], [1.25, 5.25], 9.1)
doc.add_picture(NOTE, width=Inches(6.2))
doc.inline_shapes[-1]._inline.docPr.set("descr", "Agile, MSA, 이해관계자 가치, 서비스 분리 기준을 정리한 수업 서브노트")
caption("그림 2. 수업 중 작성한 서브노트 스냅샷")

page_break()
doc.add_heading("5.2 참고 자료", level=2)
for x in [
    "최승우, 「Day1 서브노트」 및 「09.02.md」, 2026-09-02.",
    "「메디브릿지 MSA 프로젝트 기획안」, 판교 4반 2조.",
    "「Agile_MSA 실습 가이드」 및 「코드 템플릿 설명 문서」, 강의 실습자료.",
    "msa-lecture Docker Compose 및 Recommend Service(FastAPI) 소스 코드.",
    "「Cloud Agile 방법론 및 MSA 개발」 교재, 임성열.",
]: bullet(x)
doc.add_heading("5.3 구현·검증 증빙 체크리스트", level=2)
table(["증빙", "확인 내용", "상태"], [
    ("Eureka 화면", "user/course/enrollment/payment/recommend 서비스 등록", "실습 시 캡처"),
    ("Frontend", "로그인, 제품 조회, 발주 신청 화면", "실습 시 캡처"),
    ("REST API", "PENDING 발주 생성 및 조회 응답", "실습 시 캡처"),
    ("Kafka 로그", "payment.completed 발행·소비", "실습 시 캡처"),
    ("상태 변화", "결제 전 PENDING → 결제 후 ACTIVE", "실습 시 캡처"),
    ("추천 결과", "기존/신규 거래처별 다른 규칙과 추천 이유", "실습 시 캡처"),
], [1.35, 3.9, 1.25], 9.0)
callout("최종 정리", "Agile은 무엇을 언제 검증할지 결정하고, MSA는 그 검증과 변경을 작은 단위로 실행하게 한다. 메디브릿지는 규칙 기반 추천으로 시작해 실제 거래처 피드백과 구매 데이터가 쌓일 때만 다음 수준의 AI로 확장한다.")

doc.core_properties.title = "Agile+MSA 과제 - 메디브릿지"
doc.core_properties.subject = "이해관계자 가치, AI 솔루션, 아키텍처, 트러블 슈팅, 레퍼런스"
doc.core_properties.author = "최승우 (판교 4반 P131)"
doc.core_properties.keywords = "Agile, MSA, MVP, AI, 메디브릿지"
doc.save(OUT)
print(OUT)
